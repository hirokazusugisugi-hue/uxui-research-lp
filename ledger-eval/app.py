#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Ledger Eval — 総勘定元帳から担当者別実績を集計し、従業員評価レポート(Excel)を生成するツール

入力 : 総勘定元帳の Excel(.xlsx) / PDF(表形式)。複数ファイル可。列構成は自動判別。
紐づけ: 補助科目・部門欄の担当者名、または摘要欄に含まれる担当者名(名簿があれば別名も対応)
出力 : 新規Excelレポート(サマリー/月次推移/明細/未割当)
形態 : GUI(tkinter)と CLI の両対応。データは端末内で完結し、外部送信しない。

使い方(CLI):
  python app.py 元帳1.xlsx 元帳2.pdf -o 評価レポート.xlsx [--names 名簿.txt] [--year 2026]
使い方(GUI):
  python app.py            ← 引数なしで起動
"""
import argparse
import datetime as dt
import os
import re
import sys
import unicodedata
from dataclasses import dataclass, field

# ============================================================
# 設定(必要に応じて編集してください)
# ============================================================

# 評価スコアの重み(合計1.0)
WEIGHTS = {
    "sales":    0.50,   # 売上高
    "profit":   0.30,   # 利益貢献(売上 - 原価 - 経費)
    "activity": 0.20,   # 取引件数(活動量)
}

# スコアによる評価ランク
RANKS = [(80, "S"), (60, "A"), (40, "B"), (20, "C"), (0, "D")]

# 勘定科目の分類キーワード
REVENUE_KEYS = ["売上", "営業収益", "販売収入", "受取手数料", "雑収入"]
COGS_KEYS    = ["仕入", "原価", "材料"]
EXPENSE_KEYS = ["費", "給料", "給与", "賃金", "賞与", "交際", "会議", "消耗", "広告",
                "宣伝", "通信", "水道", "光熱", "地代", "家賃", "手数料", "外注",
                "荷造", "運賃", "修繕", "保険", "租税", "公課", "雑損", "燃料", "リース"]

# 列見出しの自動判別キーワード
HEADER_MAP = {
    "date":    ["日付", "年月日", "取引日", "伝票日付"],
    "account": ["勘定科目", "科目"],
    "sub":     ["補助科目", "補助"],
    "dept":    ["部門", "部署"],
    "memo":    ["摘要", "適用", "取引内容", "内容", "備考"],
    "debit":   ["借方"],
    "credit":  ["貸方"],
    "amount":  ["金額", "取引金額"],
}

# 担当者として扱わない補助科目・部門の値(銀行口座や共通部門など)
NOT_PERSON_KEYS = ["銀行", "信金", "信用金庫", "普通", "当座", "共通", "全社", "本社",
                   "その他", "一般", "口座", "支店", "現金", "小口"]

# レポート配色(UX/UI研究会パレット)
COL_ACCENT = "4F46E5"   # インディゴ
COL_SKY    = "0EA5E9"
COL_EM     = "059669"
COL_INK    = "0F172A"
COL_LIGHT  = "EEF2FF"

# ============================================================
# データ構造
# ============================================================

@dataclass
class Entry:
    date: object = None          # datetime.date or None
    account: str = ""
    sub: str = ""
    dept: str = ""
    memo: str = ""
    debit: float = 0.0
    credit: float = 0.0
    source: str = ""
    employee: str = ""
    kind: str = "other"          # revenue / cogs / expense / other

    @property
    def month(self):
        return self.date.strftime("%Y-%m") if self.date else "日付不明"


@dataclass
class Stats:
    sales: float = 0.0
    cogs: float = 0.0
    expense: float = 0.0
    count: int = 0
    monthly_sales: dict = field(default_factory=dict)

    @property
    def gross(self):   # 粗利
        return self.sales - self.cogs

    @property
    def profit(self):  # 利益貢献
        return self.gross - self.expense


# ============================================================
# 共通ユーティリティ
# ============================================================

def norm_text(s):
    if s is None:
        return ""
    return unicodedata.normalize("NFKC", str(s)).strip()


def parse_amount(v):
    """金額セルを float に。カンマ・円記号・全角・()負数に対応。"""
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = norm_text(v).replace(",", "").replace("¥", "").replace("円", "").replace(" ", "")
    if not s or s in ("-", "―", "—"):
        return 0.0
    neg = (s.startswith("(") and s.endswith(")")) or s.startswith("△") or s.startswith("▲")
    s = s.strip("()△▲")
    try:
        return -float(s) if neg else float(s)
    except ValueError:
        return 0.0


WAREKI = {"令和": 2018, "R": 2018, "平成": 1988, "H": 1988, "昭和": 1925, "S": 1925}

def parse_date(v, default_year=None):
    """日付セルを date に。datetime / 'YYYY/M/D' / '令和6年4月1日' / 'R6.4.1' / 'M/D' に対応。"""
    if v is None:
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = norm_text(v)
    if not s:
        return None
    m = re.match(r"(令和|平成|昭和|[RHS])\s*(\d+)[年.](\d+)[月.](\d+)", s)
    if m:
        base = WAREKI[m.group(1)]
        try:
            return dt.date(base + int(m.group(2)), int(m.group(3)), int(m.group(4)))
        except ValueError:
            return None
    m = re.match(r"(\d{4})[/\-年.](\d{1,2})[/\-月.](\d{1,2})", s)
    if m:
        try:
            return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            return None
    m = re.match(r"(\d{1,2})[/\-月.](\d{1,2})", s)
    if m and default_year:
        try:
            return dt.date(default_year, int(m.group(1)), int(m.group(2)))
        except ValueError:
            return None
    return None


def classify_account(name):
    n = norm_text(name)
    for k in REVENUE_KEYS:
        if k in n:
            return "revenue"
    for k in COGS_KEYS:
        if k in n:
            return "cogs"
    for k in EXPENSE_KEYS:
        if k in n:
            return "expense"
    return "other"


def detect_header(rows, max_scan=20):
    """先頭付近の行から列見出し行を探し、(行index, {field: 列index}) を返す。"""
    best = (None, {})
    for i, row in enumerate(rows[:max_scan]):
        cols = {}
        for j, cell in enumerate(row):
            t = norm_text(cell)
            if not t:
                continue
            for f, keys in HEADER_MAP.items():
                if f not in cols and any(k in t for k in keys):
                    cols[f] = j
        # 金額系(借方/貸方 or 金額)を含む行を見出しとみなす
        has_amt = ("debit" in cols or "credit" in cols or "amount" in cols)
        if has_amt and len(cols) > len(best[1]):
            best = (i, cols)
    return best


def rows_to_entries(rows, cols, header_idx, source, default_account="", default_year=None):
    """見出し行より下のデータ行を Entry に変換。"""
    entries = []
    last_date = None
    for row in rows[header_idx + 1:]:
        def cell(f):
            j = cols.get(f)
            return row[j] if j is not None and j < len(row) else None
        account = norm_text(cell("account")) or default_account
        memo = norm_text(cell("memo"))
        d = parse_date(cell("date"), default_year)
        if d:
            last_date = d
        elif memo:            # 日付セルが空欄でも直前の日付を引き継ぐ(元帳でよくある形)
            d = last_date
        debit = parse_amount(cell("debit"))
        credit = parse_amount(cell("credit"))
        if "amount" in cols and not debit and not credit:
            debit = parse_amount(cell("amount"))
        # 合計・繰越行や空行はスキップ
        joined = account + memo
        if not debit and not credit:
            continue
        if any(k in joined for k in ("合計", "小計", "繰越", "残高")):
            continue
        entries.append(Entry(
            date=d, account=account,
            sub=norm_text(cell("sub")), dept=norm_text(cell("dept")),
            memo=memo, debit=debit, credit=credit, source=source,
        ))
    return entries


# ============================================================
# 読み込み: Excel
# ============================================================

def read_excel(path, default_year=None, log=print):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    entries = []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        if not rows:
            continue
        header_idx, cols = detect_header(rows)
        if header_idx is None:
            log(f"  ⚠ シート「{ws.title}」: 見出し行を検出できずスキップ")
            continue
        # 勘定科目列がない場合、見出しより上の行 or シート名から科目名を推定
        default_account = ""
        if "account" not in cols:
            for r in rows[:header_idx]:
                for c in r:
                    m = re.search(r"勘定科目[:：]?\s*(\S+)", norm_text(c))
                    if m:
                        default_account = m.group(1)
            if not default_account:
                default_account = ws.title
        got = rows_to_entries(rows, cols, header_idx, f"{os.path.basename(path)}:{ws.title}",
                              default_account, default_year)
        log(f"  シート「{ws.title}」: {len(got)}行 取込")
        entries += got
    return entries


# ============================================================
# 読み込み: PDF
# ============================================================

def read_pdf(path, default_year=None, log=print):
    import pdfplumber
    entries = []
    with pdfplumber.open(path) as pdf:
        for pno, page in enumerate(pdf.pages, 1):
            # ページ上部のテキストから科目名を推定
            default_account = ""
            text = page.extract_text() or ""
            m = re.search(r"勘定科目[:：]?\s*(\S+)", text)
            if m:
                default_account = m.group(1)
            tables = page.extract_tables() or []
            got_n = 0
            for tbl in tables:
                rows = [list(r) for r in tbl if r]
                header_idx, cols = detect_header(rows)
                if header_idx is None:
                    continue
                got = rows_to_entries(rows, cols, header_idx,
                                      f"{os.path.basename(path)}:p{pno}",
                                      default_account, default_year)
                entries += got
                got_n += len(got)
            log(f"  {pno}ページ: {got_n}行 取込")
    if not entries:
        log("  ⚠ PDFから表を抽出できませんでした。画像スキャンのPDF(文字情報なし)は非対応です。")
    return entries


# ============================================================
# 担当者の紐づけ
# ============================================================

def load_names(path):
    """名簿ファイル: 1行1名。「正式名,別名1,別名2」形式で別名も定義可。"""
    people = {}   # alias(normalized) -> 正式名
    with open(path, encoding="utf-8-sig") as f:
        for line in f:
            parts = [norm_text(p) for p in re.split(r"[,、\t]", line) if norm_text(p)]
            if not parts:
                continue
            for alias in parts:
                people[alias] = parts[0]
    return people


def looks_like_person(value):
    v = norm_text(value)
    if not v or len(v) > 10:
        return False
    return not any(k in v for k in NOT_PERSON_KEYS)


def assign_employees(entries, people=None, log=print):
    """補助科目 → 部門 → 摘要 の順で担当者を判定する。"""
    if people is None:
        # 自動モード: 補助科目・部門に現れる「人名らしい値」を担当者候補にする
        people = {}
        for e in entries:
            for v in (e.sub, e.dept):
                if looks_like_person(v):
                    people.setdefault(norm_text(v), norm_text(v))
        log(f"  名簿なし → 補助科目・部門から担当者を自動検出: {len(set(people.values()))}名")
    aliases = sorted(people.keys(), key=len, reverse=True)   # 長い名前から先に照合
    for e in entries:
        e.kind = classify_account(e.account)
        emp = ""
        for v in (e.sub, e.dept):
            v = norm_text(v)
            if v in people:
                emp = people[v]
                break
        if not emp:
            hay = norm_text(e.sub) + " " + norm_text(e.dept) + " " + norm_text(e.memo)
            for a in aliases:
                if a and a in hay:
                    emp = people[a]
                    break
        e.employee = emp
    return people


def aggregate(entries):
    stats = {}
    for e in entries:
        if not e.employee:
            continue
        st = stats.setdefault(e.employee, Stats())
        st.count += 1
        if e.kind == "revenue":
            amt = e.credit - e.debit
            st.sales += amt
            st.monthly_sales[e.month] = st.monthly_sales.get(e.month, 0) + amt
        elif e.kind == "cogs":
            st.cogs += e.debit - e.credit
        elif e.kind == "expense":
            st.expense += e.debit - e.credit
    return stats


def score_of(st, maxima):
    def nz(x, mx):
        return max(0.0, x) / mx if mx > 0 else 0.0
    s = (WEIGHTS["sales"] * nz(st.sales, maxima["sales"])
         + WEIGHTS["profit"] * nz(st.profit, maxima["profit"])
         + WEIGHTS["activity"] * nz(st.count, maxima["activity"]))
    return round(s * 100, 1)


def rank_of(score):
    for th, r in RANKS:
        if score >= th:
            return r
    return "D"


# ============================================================
# レポート出力(Excel)
# ============================================================

def write_report(out_path, entries, stats, inputs, log=print):
    import openpyxl
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    head_fill = PatternFill("solid", fgColor=COL_ACCENT)
    head_font = Font(color="FFFFFF", bold=True)
    title_font = Font(size=14, bold=True, color=COL_INK)
    band_fill = PatternFill("solid", fgColor=COL_LIGHT)
    yen = '#,##0'

    def style_header(ws, row, ncols):
        for j in range(1, ncols + 1):
            c = ws.cell(row=row, column=j)
            c.fill = head_fill
            c.font = head_font
            c.alignment = Alignment(horizontal="center")

    # ---------- サマリー ----------
    ws = wb.active
    ws.title = "サマリー"
    ws["A1"] = "担当者別 実績集計・評価レポート"
    ws["A1"].font = title_font
    months = sorted({e.month for e in entries if e.employee and e.month != "日付不明"})
    period = f"{months[0]} 〜 {months[-1]}" if months else "-"
    ws["A2"] = f"対象期間: {period}   入力: {' / '.join(os.path.basename(p) for p in inputs)}"
    ws["A3"] = (f"スコア = 売上{int(WEIGHTS['sales']*100)}% + 利益貢献{int(WEIGHTS['profit']*100)}%"
                f" + 活動件数{int(WEIGHTS['activity']*100)}%(各項目は最大値=100で正規化)")
    ws["A2"].font = ws["A3"].font = Font(size=10, color="64748B")

    headers = ["順位", "担当者", "売上高", "売上原価", "粗利", "経費", "利益貢献", "件数", "スコア", "評価"]
    HR = 5
    for j, h in enumerate(headers, 1):
        ws.cell(row=HR, column=j, value=h)
    style_header(ws, HR, len(headers))

    maxima = {
        "sales": max((s.sales for s in stats.values()), default=0),
        "profit": max((s.profit for s in stats.values()), default=0),
        "activity": max((s.count for s in stats.values()), default=0),
    }
    order = sorted(stats.items(), key=lambda kv: score_of(kv[1], maxima), reverse=True)
    for i, (name, st) in enumerate(order, 1):
        sc = score_of(st, maxima)
        row = [i, name, round(st.sales), round(st.cogs), round(st.gross),
               round(st.expense), round(st.profit), st.count, sc, rank_of(sc)]
        for j, v in enumerate(row, 1):
            c = ws.cell(row=HR + i, column=j, value=v)
            if j in (3, 4, 5, 6, 7):
                c.number_format = yen
            if i % 2 == 0:
                c.fill = band_fill
        ws.cell(row=HR + i, column=10).font = Font(bold=True, color=COL_ACCENT)
    widths = [6, 14, 13, 13, 13, 13, 13, 7, 8, 6]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w

    if order:
        ch = BarChart()
        ch.title = "担当者別 売上高"
        ch.height, ch.width = 8, 18
        data = Reference(ws, min_col=3, min_row=HR, max_row=HR + len(order))
        cats = Reference(ws, min_col=2, min_row=HR + 1, max_row=HR + len(order))
        ch.add_data(data, titles_from_data=True)
        ch.set_categories(cats)
        ch.legend = None
        ws.add_chart(ch, f"A{HR + len(order) + 3}")

    # ---------- 月次推移 ----------
    ws2 = wb.create_sheet("月次推移")
    ws2["A1"] = "担当者別 月次売上推移"
    ws2["A1"].font = title_font
    ws2.cell(row=3, column=1, value="担当者")
    for j, m in enumerate(months, 2):
        ws2.cell(row=3, column=j, value=m)
    ws2.cell(row=3, column=len(months) + 2, value="合計")
    style_header(ws2, 3, len(months) + 2)
    for i, (name, st) in enumerate(order, 1):
        ws2.cell(row=3 + i, column=1, value=name)
        for j, m in enumerate(months, 2):
            c = ws2.cell(row=3 + i, column=j, value=round(st.monthly_sales.get(m, 0)))
            c.number_format = yen
        t = ws2.cell(row=3 + i, column=len(months) + 2, value=round(st.sales))
        t.number_format = yen
        t.font = Font(bold=True)
    ws2.column_dimensions["A"].width = 14
    for j in range(2, len(months) + 3):
        ws2.column_dimensions[get_column_letter(j)].width = 11
    if order and months:
        ch2 = LineChart()
        ch2.title = "月次売上推移"
        ch2.height, ch2.width = 9, 20
        data = Reference(ws2, min_col=1, min_row=3, max_col=len(months) + 1,
                         max_row=3 + len(order))
        ch2.add_data(data, titles_from_data=True, from_rows=True)
        ws2.add_chart(ch2, f"A{len(order) + 6}")

    # ---------- 明細 / 未割当 ----------
    def detail_sheet(title, rows_src, with_emp):
        w = wb.create_sheet(title)
        cols = (["担当者"] if with_emp else []) + \
               ["日付", "勘定科目", "区分", "補助科目", "部門", "摘要", "借方", "貸方", "出典"]
        for j, h in enumerate(cols, 1):
            w.cell(row=1, column=j, value=h)
        style_header(w, 1, len(cols))
        kind_jp = {"revenue": "売上", "cogs": "原価", "expense": "経費", "other": "その他"}
        for i, e in enumerate(rows_src, 2):
            vals = ([e.employee] if with_emp else []) + \
                   [e.date.isoformat() if e.date else "", e.account, kind_jp[e.kind],
                    e.sub, e.dept, e.memo, e.debit or None, e.credit or None, e.source]
            for j, v in enumerate(vals, 1):
                c = w.cell(row=i, column=j, value=v)
                if cols[j - 1] in ("借方", "貸方"):
                    c.number_format = yen
        for j, wd in enumerate(([12] if with_emp else []) + [11, 16, 7, 12, 12, 34, 12, 12, 22], 1):
            w.column_dimensions[get_column_letter(j)].width = wd
        w.freeze_panes = "A2"

    matched = sorted([e for e in entries if e.employee],
                     key=lambda e: (e.employee, e.date or dt.date.min))
    unmatched = [e for e in entries if not e.employee and e.kind != "other"]
    detail_sheet("明細(担当者別)", matched, with_emp=True)
    detail_sheet("未割当", unmatched, with_emp=False)

    wb.save(out_path)
    log(f"✅ レポートを出力しました: {out_path}")
    log(f"   担当者 {len(stats)}名 / 紐づけ済 {len(matched)}行 / 未割当(損益系) {len(unmatched)}行")


# ============================================================
# 実行パイプライン
# ============================================================

def run(inputs, output, names_path=None, default_year=None, log=print):
    entries = []
    for p in inputs:
        log(f"📖 読み込み: {p}")
        ext = os.path.splitext(p)[1].lower()
        if ext in (".xlsx", ".xlsm"):
            entries += read_excel(p, default_year, log)
        elif ext == ".pdf":
            entries += read_pdf(p, default_year, log)
        else:
            log(f"  ⚠ 未対応の形式です: {ext}(xlsx / pdf のみ)")
    if not entries:
        log("❌ 取り込めた仕訳が0行でした。見出し行(日付・摘要・借方・貸方など)があるか確認してください。")
        return False
    log(f"🧾 仕訳 {len(entries)}行 取込完了")
    people = load_names(names_path) if names_path else None
    if people:
        log(f"👥 名簿: {len(set(people.values()))}名(別名含め{len(people)}表記)")
    assign_employees(entries, people, log)
    stats = aggregate(entries)
    if not stats:
        log("❌ 担当者を1名も紐づけできませんでした。名簿(--names)の指定、または補助科目/部門/摘要への担当者名の記載を確認してください。")
        return False
    write_report(output, entries, stats, inputs, log)
    return True


# ============================================================
# GUI(tkinter)
# ============================================================

def launch_gui():
    try:
        import tkinter as tk
        from tkinter import filedialog, messagebox, scrolledtext
    except ImportError:
        print("tkinter が見つかりません。CLIで実行してください:")
        print("  python app.py 元帳.xlsx -o レポート.xlsx")
        sys.exit(1)

    root = tk.Tk()
    root.title("Ledger Eval — 担当者実績集計・評価")
    root.geometry("640x520")

    files = []

    frm = tk.Frame(root, padx=12, pady=10)
    frm.pack(fill="both", expand=True)

    tk.Label(frm, text="① 総勘定元帳ファイル(Excel / PDF、複数可)", anchor="w").pack(fill="x")
    lst = tk.Listbox(frm, height=5)
    lst.pack(fill="x")
    row1 = tk.Frame(frm); row1.pack(fill="x", pady=4)

    def add_files():
        for p in filedialog.askopenfilenames(
                filetypes=[("元帳ファイル", "*.xlsx *.xlsm *.pdf")]):
            if p not in files:
                files.append(p)
                lst.insert("end", p)

    def clear_files():
        files.clear(); lst.delete(0, "end")

    tk.Button(row1, text="ファイルを追加", command=add_files).pack(side="left")
    tk.Button(row1, text="クリア", command=clear_files).pack(side="left", padx=6)

    tk.Label(frm, text="② 担当者名簿(任意・1行1名、「正式名,別名」も可。未指定なら補助科目/部門から自動検出)",
             anchor="w").pack(fill="x", pady=(8, 0))
    row2 = tk.Frame(frm); row2.pack(fill="x")
    names_var = tk.StringVar()
    tk.Entry(row2, textvariable=names_var).pack(side="left", fill="x", expand=True)
    tk.Button(row2, text="選択",
              command=lambda: names_var.set(filedialog.askopenfilename(
                  filetypes=[("名簿", "*.txt *.csv")]) or names_var.get())).pack(side="left", padx=6)

    row3 = tk.Frame(frm); row3.pack(fill="x", pady=(8, 0))
    tk.Label(row3, text="③ 日付に年がない場合の既定年:").pack(side="left")
    year_var = tk.StringVar(value=str(dt.date.today().year))
    tk.Entry(row3, textvariable=year_var, width=8).pack(side="left", padx=6)

    logbox = scrolledtext.ScrolledText(frm, height=12, state="disabled")

    def log(msg):
        logbox.configure(state="normal")
        logbox.insert("end", str(msg) + "\n")
        logbox.see("end")
        logbox.configure(state="disabled")
        root.update_idletasks()

    def execute():
        if not files:
            messagebox.showwarning("Ledger Eval", "元帳ファイルを追加してください")
            return
        out = filedialog.asksaveasfilename(
            defaultextension=".xlsx", initialfile="評価レポート.xlsx",
            filetypes=[("Excel", "*.xlsx")])
        if not out:
            return
        try:
            year = int(year_var.get()) if year_var.get().strip() else None
        except ValueError:
            year = None
        try:
            ok = run(files, out, names_var.get() or None, year, log)
            if ok:
                messagebox.showinfo("Ledger Eval", f"完了しました:\n{out}")
        except Exception as ex:
            log(f"❌ エラー: {ex}")
            messagebox.showerror("Ledger Eval", str(ex))

    tk.Button(frm, text="④ レポート生成", bg="#4F46E5", fg="white",
              command=execute).pack(fill="x", pady=10)
    logbox.pack(fill="both", expand=True)
    log("元帳ファイルを追加して「レポート生成」を押してください。データは外部送信されません。")
    root.mainloop()


# ============================================================
# エントリポイント
# ============================================================

def main():
    if len(sys.argv) == 1:
        launch_gui()
        return
    ap = argparse.ArgumentParser(description="総勘定元帳 → 担当者別実績集計・評価レポート(Excel)")
    ap.add_argument("inputs", nargs="+", help="元帳ファイル(.xlsx / .pdf)")
    ap.add_argument("-o", "--output", default="評価レポート.xlsx", help="出力Excelパス")
    ap.add_argument("--names", help="担当者名簿ファイル(1行1名、「正式名,別名」可)")
    ap.add_argument("--year", type=int, help="日付に年がない場合に補う年(例: 2026)")
    args = ap.parse_args()
    ok = run(args.inputs, args.output, args.names, args.year)
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
